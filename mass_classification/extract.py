"""Bounded local file extraction. No uploaded content is executed or fetched."""
import csv
import email
from email import policy
from html.parser import HTMLParser
import io
import json
import mimetypes
from pathlib import Path
import re
import subprocess
import tempfile
import zipfile


class ExtractionError(ValueError):
    pass


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.hidden += 1

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def _run(args: list[str], timeout: int = 90) -> bytes:
    try:
        result = subprocess.run(args, capture_output=True, check=True, timeout=timeout)
        return result.stdout
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, FileNotFoundError) as exc:
        raise ExtractionError(f"Conversion failed: {args[0]}") from exc


def _transcribe(path: Path, model_name: str) -> str:
    from faster_whisper import WhisperModel
    model = WhisperModel(model_name, device="cpu", compute_type="int8")
    segments, _ = model.transcribe(str(path), vad_filter=True)
    return " ".join(segment.text for segment in segments)


def extract(path: Path, filename: str, asr_model: str = "small", depth: int = 0) -> tuple[str, dict]:
    """Return extracted text and metadata; raise on unsupported or corrupt input."""
    if depth > 2:
        raise ExtractionError("Archive nesting exceeds limit")
    suffix = Path(filename).suffix.lower()
    raw_head = path.open("rb").read(16)
    meta: dict = {"format": suffix, "bytes": path.stat().st_size}
    if raw_head.startswith(b"%PDF"):
        import fitz
        from PIL import Image
        import pytesseract
        with fitz.open(path) as pdf:
            if len(pdf) > 1000:
                raise ExtractionError("PDF page limit exceeded")
            pages = []
            for page in pdf:
                content = page.get_text()
                if len(content.strip()) < 20:
                    pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
                    content += pytesseract.image_to_string(Image.open(io.BytesIO(pix.tobytes("png"))))
                pages.append(content)
            meta["pages"] = len(pdf)
            text = "\n".join(pages)
    elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}:
        from PIL import Image, ExifTags
        import pytesseract
        with Image.open(path) as img:
            img.verify()
        with Image.open(path) as img:
            if img.width * img.height > 80_000_000:
                raise ExtractionError("Image pixel limit exceeded")
            meta["dimensions"] = [img.width, img.height]
            exif = img.getexif()
            meta["exif"] = {str(ExifTags.TAGS.get(k, k)): str(v)[:256] for k, v in exif.items()
                                if ExifTags.TAGS.get(k, k) not in ("GPSInfo", "MakerNote")}
            text = pytesseract.image_to_string(img)
    elif suffix == ".docx":
        from docx import Document
        document = Document(path)
        text = "\n".join([p.text for p in document.paragraphs] +
                         [" | ".join(c.text for c in r.cells) for t in document.tables for r in t.rows])
        meta["author"] = document.core_properties.author
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            text = "\n".join(" | ".join(str(c) if c is not None else "" for c in row)
                             for sheet in book for row in sheet.iter_rows(values_only=True))
            meta["sheets"] = book.sheetnames
        finally:
            book.close()
    elif suffix in {".mp3", ".wav", ".ogg", ".flac", ".m4a"}:
        text = _transcribe(path, asr_model)
    elif suffix in {".mp4", ".mkv", ".webm", ".mov"}:
        from PIL import Image
        import pytesseract
        with tempfile.TemporaryDirectory() as tmp:
            frames = Path(tmp) / "frame-%03d.jpg"
            _run(["ffmpeg", "-v", "error", "-i", str(path), "-vf", "fps=1/30,scale=960:-1", "-frames:v", "120", str(frames), "-y"])
            captions = [pytesseract.image_to_string(Image.open(frame)) for frame in sorted(Path(tmp).glob("frame-*.jpg"))]
            audio = Path(tmp) / "audio.wav"
            try:
                _run(["ffmpeg", "-v", "error", "-i", str(path), "-vn", "-ac", "1", "-ar", "16000", str(audio), "-y"])
                spoken = _transcribe(audio, asr_model)
            except ExtractionError:
                spoken = ""
            text = spoken + "\n" + "\n".join(captions)
            meta["sampled_frames"] = len(captions)
    elif suffix == ".eml":
        msg = email.message_from_binary_file(path.open("rb"), policy=policy.default)
        meta.update({"subject": str(msg.get("subject", "")), "from": str(msg.get("from", "")), "date": str(msg.get("date", ""))})
        text = "\n".join(str(part.get_content()) for part in msg.walk()
                         if part.get_content_type() == "text/plain" and not part.get_filename())
    elif suffix == ".msg":
        import extract_msg
        msg = extract_msg.Message(str(path))
        meta.update({"subject": msg.subject or "", "from": msg.sender or "", "date": str(msg.date or "")})
        text = msg.body or ""
    elif suffix == ".zip":
        pieces = []
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 200 or sum(m.file_size for m in members) > 100_000_000:
                raise ExtractionError("Archive expansion limit exceeded")
            for member in members:
                if member.is_dir() or member.filename.startswith("/") or ".." in Path(member.filename).parts:
                    continue
                with tempfile.NamedTemporaryFile() as tmp:
                    tmp.write(archive.read(member)); tmp.flush()
                    try:
                        child, _ = extract(Path(tmp.name), member.filename, asr_model, depth + 1)
                        pieces.append(f"[{member.filename}]\n{child}")
                    except ExtractionError:
                        continue
        text = "\n".join(pieces)
    elif suffix in {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".yaml", ".yml", ".html", ".htm", ".py", ".js", ".sql", ".log"}:
        blob = path.read_bytes()
        if b"\x00" in blob:
            raise ExtractionError("Binary content in text file")
        try:
            text = blob.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = blob.decode("latin-1")
            meta["encoding_fallback"] = "latin-1"
        if suffix in {".html", ".htm"}:
            parser = _HTMLText(); parser.feed(text); text = " ".join(parser.parts)
        if suffix == ".json":
            json.loads(text)
        if suffix in {".csv", ".tsv"}:
            dialect = "excel-tab" if suffix == ".tsv" else "excel"
            rows = csv.reader(io.StringIO(text), dialect=dialect)
            text = "\n".join(" | ".join(row) for row in rows)
    else:
        raise ExtractionError(f"Unsupported file type: {suffix or 'unknown'}")
    text = re.sub(r"[ \t]+", " ", text).strip()[:2_000_000]
    if not text:
        raise ExtractionError("No extractable text; manual review required")
    return text, meta


def chunks(text: str, size: int = 900, overlap: int = 100) -> list[str]:
    if size <= overlap or overlap < 0:
        raise ValueError("Invalid chunk parameters")
    return [text[start:start + size] for start in range(0, len(text), size - overlap)
            if text[start:start + size].strip()]
